#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从《运维系统中的共享内存定义.xlsx》提取数据定义表，生成 services/dataplane/seed/ 下的种子 CSV。

产物：
  sensor_data_defi.csv  值型变量定义（category,name,add,std_tag,accuracy,unit,org,signed,sheet）
  status_data_defi.csv  布尔变量定义（category,status,add,std_tag,org,sheet）
  fault_map.csv         故障位映射（fault_code,org,name,trigger,module_name）
  ops_map.csv           操作组映射（group_no,source_name,source_type,addr_hi,addr_lo,org_hi,org_lo）
  fault_rule.csv / fault_boc_hpc_rule.csv / fault_moc_rule.csv  从 shearer-remote-ops 仓库照搬

ORG 列写法（详见 docs/org解析规则.md 第 5 节）：
  "2"            org 第 2 列原值
  "282Bit12"     第 282 列的 bit12（布尔）
  "156Bit2-C"    同上，故障触发逻辑反转
  "498Bit12-15"  第 498 列的 bit12~15（4 位半字节值，按值型处理）
  "291-C"+状态值  原值直写，值==状态值 触发故障
  "514-CN"+状态值 原值直写，值!=状态值 触发故障（==状态值 时消除）
  "ecspds.512.0-C" 外部信号写法，按 512Bit0 反逻辑布尔处理（Excel 疑点，见文档）
  空白            无 org 数据源，恒写 0
"""
import csv, os, re, shutil, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "运维系统中的共享内存定义.xlsx")
SEED = os.path.join(ROOT, "services", "dataplane", "seed")
REMOTE_OPS_CSV = "/mnt/c/Users/wu_ve/Desktop/shearer-remote-ops/server/backend/internal/server"

# ---------------------------------------------------------------------------
# 地址 → std_tag 全量映射（地址在整块共享内存中唯一）
# ---------------------------------------------------------------------------
ADDR_TAGS = {
    # ===== 电机 sheet =====
    0: "left_cutter_a_current_u", 1: "left_cutter_a_current_v", 2: "left_cutter_a_current_w",
    3: "left_cutter_a_load_pct", 4: "left_cutter_a_temp",
    5: "left_cutter_b_current_u", 6: "left_cutter_b_current_v", 7: "left_cutter_b_current_w",
    8: "left_cutter_b_load_pct", 9: "left_cutter_b_temp",
    10: "right_cutter_a_current_u", 11: "right_cutter_a_current_v", 12: "right_cutter_a_current_w",
    13: "right_cutter_a_load_pct", 14: "right_cutter_a_temp",
    15: "right_cutter_b_current_u", 16: "right_cutter_b_current_v", 17: "right_cutter_b_current_w",
    18: "right_cutter_b_load_pct", 19: "right_cutter_b_temp",
    20: "left_haul_current_u", 21: "left_haul_current_v", 22: "left_haul_current_w",
    23: "left_haul_load_pct", 24: "left_haul_temp",
    25: "right_haul_current_u", 26: "right_haul_current_v", 27: "right_haul_current_w",
    28: "right_haul_load_pct", 29: "right_haul_temp",
    30: "left_pump_current_u", 31: "left_pump_current_v", 32: "left_pump_current_w",
    33: "left_pump_load_pct", 34: "left_pump_temp",
    35: "right_pump_current_u", 36: "right_pump_current_v", 37: "right_pump_current_w",
    38: "right_pump_load_pct", 39: "right_pump_temp",
    40: "crusher_current_u", 41: "crusher_current_v", 42: "crusher_current_w",
    43: "crusher_load_pct", 44: "crusher_temp",
    45: "haul_xfmr_current_u", 46: "haul_xfmr_current_v", 47: "haul_xfmr_current_w",
    48: "haul_xfmr_load_pct", 49: "haul_xfmr_temp",
    50: "left_cutter_start_warn", 51: "left_cutter_starting", 52: "left_cutter_contactor_on",
    53: "left_cutter_running", 54: "left_cutter_leak_lock",
    55: "right_cutter_start_warn", 56: "right_cutter_starting", 57: "right_cutter_contactor_on",
    58: "right_cutter_running", 59: "right_cutter_leak_lock",
    60: "pump_start_warn", 61: "pump_starting", 62: "pump_contactor_on",
    63: "pump_running", 64: "pump_leak_lock",
    65: "crusher_start_warn", 66: "crusher_starting", 67: "crusher_contactor_on",
    68: "crusher_running", 69: "crusher_leak_lock",
    80: "left_cutter_vfd_current", 81: "right_cutter_vfd_current",
    # ===== 变频器 sheet =====
    600: "left_inv_input_voltage_a", 601: "left_inv_input_voltage_b", 602: "left_inv_input_voltage_c",
    603: "left_inv_control_word",
    604: "left_inv_fwd_run", 605: "left_inv_rev_run", 606: "left_inv_ext_fault", 607: "left_inv_fault_reset",
    621: "left_inv_output_freq", 622: "left_inv_output_torque", 623: "left_inv_output_current",
    624: "left_inv_bus_voltage", 625: "left_inv_igbt_temp",
    661: "right_inv_output_freq", 662: "right_inv_output_torque", 663: "right_inv_output_current",
    664: "right_inv_bus_voltage", 665: "right_inv_igbt_temp",
    # ===== 传感器 sheet =====
    822: "left_pumpbox_temp", 823: "left_haulbox_temp", 824: "left_arm_temp",
    825: "left_haulbox_oil_level", 826: "left_pumpbox_oil_level",
    827: "left_swing_sensor_ok", 828: "left_arm_height", 829: "left_arm_swing_angle",
    830: "left_swing_raw", 831: "left_water_pressure", 832: "left_back_pressure",
    833: "left_arm_oil_level", 834: "left_cable_tension_ok", 835: "left_pos_encoder_ok",
    836: "left_speed_measured", 837: "left_shield_no", 838: "left_position",
    839: "left_encoder_raw1", 840: "left_encoder_raw2", 841: "left_sync_switch",
    842: "left_water_flow1", 843: "left_water_flow2", 844: "left_arm_flow",
    852: "right_pumpbox_temp", 853: "right_haulbox_temp", 854: "right_arm_temp",
    855: "right_haulbox_oil_level", 856: "right_pumpbox_oil_level",
    857: "right_swing_sensor_ok", 858: "right_arm_height", 859: "right_arm_swing_angle",
    860: "right_swing_raw", 861: "right_water_pressure", 862: "right_back_pressure",
    863: "right_arm_oil_level", 864: "right_cable_tension_ok", 865: "right_pos_encoder_ok",
    866: "right_speed_measured", 867: "right_shield_no", 868: "right_position",
    869: "right_encoder_raw1", 870: "right_encoder_raw2", 871: "right_sync_switch",
    872: "right_water_flow1", 873: "right_water_flow2", 874: "right_arm_flow",
    885: "left_aux_cyl_stroke_pos", 886: "left_aux_cyl_stroke_speed", 887: "left_aux_cyl_arm_angle",
    888: "left_aux_cyl_arm_angle_hp_l", 889: "left_aux_cyl_arm_angle_hp_h",
    890: "left_cyl_conn_ok", 891: "left_cyl_push_stroke", 892: "left_cyl_arm_angle",
    895: "right_aux_cyl_stroke_pos", 896: "right_aux_cyl_stroke_speed", 897: "right_aux_cyl_arm_angle",
    898: "right_aux_cyl_arm_angle_hp_l", 899: "right_aux_cyl_arm_angle_hp_h",
    900: "right_cyl_conn_ok", 901: "right_cyl_push_stroke", 902: "right_cyl_arm_angle",
    905: "dust_density",
    910: "sps_encoder_speed", 911: "sps_encoder_pos", 912: "sps_encoder_pos_hp_l", 913: "sps_encoder_pos_hp_h",
    919: "sensor_status_word",
    920: "left_radar_ok", 921: "right_radar_ok", 922: "left_radar_collision_warn", 923: "right_radar_collision_warn",
    924: "left_radar_status_word", 925: "left_radar_avg_dist", 926: "left_radar_inst_dist", 927: "left_radar_signal",
    929: "right_radar_status_word", 930: "right_radar_avg_dist", 931: "right_radar_inst_dist", 932: "right_radar_signal",
    934: "sensor_status_word1",
    935: "left_arm_incl_comm_state", 936: "ecbox_incl_comm_state", 937: "left_stroke_comm_state",
    938: "comm_state_a938",
    939: "sensor_status_word2",
    940: "right_arm_incl_comm_state", 941: "comm_state_a941", 942: "right_stroke_comm_state",
    943: "sps_encoder_comm_state",
    944: "left_arm_incl_x", 945: "left_arm_incl_y", 946: "right_arm_incl_x", 947: "right_arm_incl_y",
    948: "ecbox_incl_x", 949: "ecbox_incl_y",
    950: "left_pds_valid", 951: "left_pds_run_state",
    952: "left_pds_comm_lost", 953: "pds_person_near_warn", 954: "pds_person_enter_protect",
    955: "dust_comm_ok", 956: "dust_high_warn", 957: "dust_high_protect",
    958: "left_pds_comm_ok", 959: "right_pds_comm_ok",
    960: "pds3_left_rx_ok", 961: "pds3_right_rx_ok", 962: "pds3_mid_rx_ok",
    # ===== 状态 sheet =====
    2000: "boc_comm_value", 2001: "hpc_comm_ok", 2002: "axi_comm_ok", 2003: "wcr_comm_ok",
    2004: "sgc_comm_ok", 2005: "moc_comm_ok", 2006: "eio_comm_ok", 2007: "lesc_comm_ok",
    2008: "resc_comm_ok", 2009: "bp1_comm_ok", 2010: "bp2_comm_ok", 2011: "lcb_comm_ok",
    2012: "rcb_comm_ok", 2013: "ebi_comm_ok", 2014: "exi_comm_ok",
    2064: "fsk_comm_active", 2065: "remote_ctrl_enabled", 2066: "gateroad_remote", 2067: "host_remote",
    2068: "lasc_enabled", 2069: "lasc_preparing", 2070: "lasc_ready",
    2096: "moving_left", 2097: "moving_right", 2098: "haul_powered", 2099: "haul_started",
    2100: "dispatch_mode", 2101: "hpc_local_ctrl", 2102: "brake_release_cmd", 2103: "brake_release_ack",
    2104: "brake_released", 2105: "cable_tow_state",
    2125: "left_cutter_locked", 2126: "right_cutter_locked",
    2127: "no_cooling_countdown", 2128: "software_unlicensed", 2129: "display_no_comm",
    2130: "error_locked_reset", 2131: "haul_stop_locked_reset", 2132: "estop_locked_reset",
    2133: "waiting_ukey_auth", 2134: "waiting_lasc_ready", 2135: "carrier_comm_fail",
    2136: "haul_reset_key_pressed",
    2137: "left_water_press_low_warn", 2138: "left_water_press_low_protect",
    2139: "right_water_press_low_warn", 2140: "right_water_press_low_protect",
    2141: "haul_not_ready_pump_blocked",
}

# 操作 sheet 的来源分组（组序号 → 来源名称/类型），地址与 ORG 从 Excel 读
OPS_GROUPS = [
    (1, "左端头站", "输入", 1902, 1903), (2, "右端头站", "输入", 1904, 1905),
    (3, "左遥控器", "输入", 1906, 1907), (4, "右遥控器", "输入", 1908, 1909),
    (5, "电控箱按钮", "输入", 1910, 1911), (6, "自动截割", "输入", 1912, 1913),
    (7, "上位机", "输入", 1914, 1915), (8, "顺槽控制站", "输入", 1916, 1917),
    (9, "系统保护", "输出", 1918, 1919), (10, "系统输出", "输出", 1920, 1921),
]

RE_BIT = re.compile(r"^(\d+)Bit(\d+)(-C)?$")            # 282Bit12 / 156Bit2-C
RE_BITRANGE = re.compile(r"^(\d+)Bit(\d+)-(\d+)$")      # 498Bit12-15
RE_VAL_C = re.compile(r"^(\d+)-(C|CN)$")                # 291-C / 514-CN
RE_ECSPDS = re.compile(r"^ecspds\.(\d+)\.(\d+)-C$")     # ecspds.512.0-C


def parse_org(org):
    """解析 ORG 写法 → (kind, col, bit_lo, bit_hi, inverted)
    kind: val / bit / bitrange / none"""
    s = str(org).strip() if org is not None else ""
    if not s or s == "None":
        return ("none", None, None, None, False)
    m = RE_BITRANGE.match(s)
    if m:
        return ("bitrange", int(m.group(1)), int(m.group(2)), int(m.group(3)), False)
    m = RE_BIT.match(s)
    if m:
        return ("bit", int(m.group(1)), int(m.group(2)), int(m.group(2)), m.group(3) == "-C")
    m = RE_ECSPDS.match(s)
    if m:  # 疑点写法：按 colBit(bit) 反逻辑布尔处理
        return ("bit", int(m.group(1)), int(m.group(2)), int(m.group(2)), True)
    m = RE_VAL_C.match(s)
    if m:
        return ("val", int(m.group(1)), None, None, True)  # -C/-CN：原值直写，触发逻辑另记
    if s.isdigit():
        return ("val", int(s), None, None, False)
    raise ValueError(f"无法解析的 ORG 写法: {org!r}")


def load_sheet(ws, sheet_name):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = next(i for i, r in enumerate(rows) if r and "地址" in [str(c) for c in r if c])
    hdr = [str(c).strip() if c is not None else "" for c in rows[hdr_i]]
    # 变频器 sheet 的 ORG 表头是空白串——按列位置回填（地址列的下一列）
    if "ORG" not in hdr:
        hdr[hdr.index("地址") + 1] = "ORG"
    items = []
    for r in rows[hdr_i + 1:]:
        if r is None or all(c is None for c in r):
            continue
        d = {hdr[j]: r[j] for j in range(min(len(hdr), len(r))) if hdr[j]}
        if d.get("地址") is None:
            continue
        d["地址"] = int(d["地址"])
        d["名称"] = str(d.get("名称") or "").strip()
        d["sheet"] = sheet_name
        items.append(d)
    return items


def main():
    os.makedirs(SEED, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    sheets = {name: load_sheet(wb[name], name) for name in wb.sheetnames}

    sensors, statuses, faults = [], [], []

    for sheet_name in ["电机", "变频器", "传感器", "状态"]:
        for d in sheets[sheet_name]:
            addr, name = d["地址"], d["名称"]
            category = str(d.get("部件") or d.get("位置") or "系统").strip()
            kind, col, blo, bhi, inv = parse_org(d.get("ORG"))
            tag = ADDR_TAGS.get(addr)
            if tag is None:
                raise SystemExit(f"[缺std_tag] {sheet_name} addr={addr} {name}")
            org_str = str(d.get("ORG") or "").strip()
            if kind == "bit" and blo == bhi:  # 单 bit → 布尔
                statuses.append(dict(category=category, status=name, add=addr,
                                     std_tag=tag, org=org_str, sheet=sheet_name))
            else:  # 值 / bit 区间 / 未映射 → 值型
                acc = d.get("精度")
                acc = float(acc) if acc is not None else 1.0
                unit = str(d.get("单位") or "").strip()
                signed = 1 if acc < 1 else 0
                sensors.append(dict(category=category, name=name, add=addr,
                                    std_tag=tag, accuracy=acc, unit=unit,
                                    org=org_str, signed=signed, sheet=sheet_name))

    # 故障 sheet → fault_map
    for d in sheets["故障"]:
        addr, name = d["地址"], d["名称"]
        sv = d.get("状态值")
        kind, col, blo, bhi, inv = parse_org(d.get("ORG"))
        org_str = str(d.get("ORG") or "").strip()
        module = ""
        if 3400 <= addr <= 3411:
            module = ["BOC", "HPC", "MOC"][(addr - 3400) // 4]
            trigger = "fifo"
        elif kind == "bit":
            trigger = "bit_inv" if inv else "bit"
        elif kind == "val" and sv is not None:
            # -C：值==状态值 触发；-CN：值==状态值 消除（其余触发）
            trigger = (f"eq:{int(sv)}" if org_str.endswith("-C") else f"neq:{int(sv)}")
        else:
            raise SystemExit(f"[故障解析失败] addr={addr} org={org_str!r} 状态值={sv}")
        faults.append(dict(fault_code=addr, org=org_str, name=name,
                           trigger=trigger, module_name=module))

    # 操作 sheet → ops_map（地址/ORG 以 Excel 为准，含疑点 1918/1919）
    ops_by_addr = {d["地址"]: d for d in sheets["操作"]}
    ops = []
    for gno, sname, stype, ahi, alo in OPS_GROUPS:
        ops.append(dict(group_no=gno, source_name=sname, source_type=stype,
                        addr_hi=ahi, addr_lo=alo,
                        org_hi=int(ops_by_addr[ahi]["ORG"]), org_lo=int(ops_by_addr[alo]["ORG"])))

    # ---- 校验 ----
    tags = [s["std_tag"] for s in sensors] + [s["std_tag"] for s in statuses]
    dup = {t for t in tags if tags.count(t) > 1}
    assert not dup, f"std_tag 重复: {dup}"
    for s in sensors + statuses:
        m = RE_BIT.match(s["org"]) or RE_BITRANGE.match(s["org"]) or RE_ECSPDS.match(s["org"])
        c = int(m.group(1)) if m else (int(s["org"]) if s["org"].isdigit() else None)
        assert c is None or 1 <= c <= 609, f"org 列越界: {s}"
    fr_codes = set()
    with open(os.path.join(REMOTE_OPS_CSV, "fault_rule_catalog.csv"), encoding="utf-8-sig") as f:
        fr_codes = {int(r["fault_code"]) for r in csv.DictReader(f)}
    fm_codes = {f_["fault_code"] for f_ in faults if f_["trigger"] != "fifo"}
    print(f"fault_map(非FIFO) {len(fm_codes)} 条, fault_rule {len(fr_codes)} 条, "
          f"map有rule无: {sorted(fm_codes - fr_codes)}, rule有map无: {sorted(fr_codes - fm_codes)}")

    # ---- 写 CSV ----
    def dump(fname, rows, fields):
        with open(os.path.join(SEED, fname), "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
        print(f"  {fname}: {len(rows)} 行")

    dump("sensor_data_defi.csv", sensors,
         ["category", "name", "add", "std_tag", "accuracy", "unit", "org", "signed", "sheet"])
    dump("status_data_defi.csv", statuses,
         ["category", "status", "add", "std_tag", "org", "sheet"])
    dump("fault_map.csv", faults, ["fault_code", "org", "name", "trigger", "module_name"])
    dump("ops_map.csv", ops,
         ["group_no", "source_name", "source_type", "addr_hi", "addr_lo", "org_hi", "org_lo"])

    # ---- 照搬三张规则表 ----
    for src, dst in [("fault_rule_catalog.csv", "fault_rule.csv"),
                     ("fault_boc_hpc_catalog.csv", "fault_boc_hpc_rule.csv"),
                     ("fault_moc_catalog.csv", "fault_moc_rule.csv")]:
        p = os.path.join(REMOTE_OPS_CSV, src)
        shutil.copy(p, os.path.join(SEED, dst))
        with open(os.path.join(SEED, dst), encoding="utf-8-sig") as f:
            n = sum(1 for _ in f) - 1
        print(f"  {dst}: {n} 行 (照搬自 {src})")

    print("完成。")


if __name__ == "__main__":
    main()
